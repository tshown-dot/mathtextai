import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

def clean_markdown_input(text: str) -> str:
    """
    Ensures markdown tables have valid dashed separator rows and proper spacing.
    """
    lines = text.strip().split("\n")
    processed = []
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        # Check if line looks like a table header (| Problem | Steps |)
        if line.startswith("|") and line.endswith("|") and ("problem" in line.lower() or "steps" in line.lower()):
            processed.append("\n" + line)
            # Check if next line is already a separator (| :--- | :--- |)
            if i + 1 < len(lines) and not re.match(r"^\|(\s*:?-+:?\s*\|)+$", lines[i + 1].strip()):
                cols = len([c for c in line.split("|")[1:-1]])
                sep = "| " + " | ".join([":---"] * cols) + " |"
                processed.append(sep)
        elif line.startswith("# PART ") or line.startswith("PART ") or line.startswith("LESSON ") or line.startswith("## Tool ") or line.startswith("Tool "):
            processed.append("\n" + line + "\n")
        else:
            processed.append(line)
        i += 1
    return "\n".join(processed)


def format_math_document(docx_path: str) -> None:
    """
    Applies professional layout, styling, and typography to math notes documents:
      1. Dynamic 2-page pagination based on tool count (3 tools vs 2 tools)
      2. Dark charcoal section banners with bold white text (#555555) for PART 1, 2, 3
      3. Slate Silver (#CBD5E1) table headers with bold text and dark grey borders (#555555)
      4. 2/3 and 1/3 column proportions for Tool tables (3.2" Problem, 1.8" Steps)
      5. Full 6.5" width for Fix the Mistake table
      6. Calibrated spacing for CYU questions to balance Page 2
    """
    try:
        doc = docx.Document(docx_path)
    except Exception as e:
        print(f"Warning: Could not open document for formatting: {e}")
        return

    w_ns = nsdecls("w")
    BORDER_COLOR = "555555"
    HEADER_BG = "CBD5E1"      # Option A: Slate Silver
    BANNER_BG = "555555"      # Dark Charcoal Banner

    # -------------------------------------------------------------
    # 1. COUNT TOOLS AND IDENTIFY SECTIONS
    # -------------------------------------------------------------
    tool_count = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if re.match(r"^Tool\s*\d+", text, re.IGNORECASE):
            tool_count += 1

    if tool_count == 0:
        tool_count = max(1, min(3, len(doc.tables) - 1))

    # Page break rule
    # 3 Tools -> Page break before Part 2
    # 2 Tools -> Page break before Part 3
    break_before_part2 = (tool_count >= 3)
    cyu_space_after = Pt(16) if break_before_part2 else Pt(28)
    cyu_space_before = Pt(4) if break_before_part2 else Pt(6)

    # -------------------------------------------------------------
    # 2. FORMAT PARAGRAPHS, BANNERS, HEADINGS, AND CYU QUESTIONS
    # -------------------------------------------------------------
    shd_banner_xml = f'<w:shd {w_ns} w:fill="{BANNER_BG}"/>'
    current_section = "INTRO"

    for p in doc.paragraphs:
        text = p.text.strip()
        text_upper = text.upper()

        if "PART 1" in text_upper:
            current_section = "PART1"
            _apply_banner_style(p, shd_banner_xml)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.page_break_before = False

        elif "PART 2" in text_upper:
            current_section = "PART2"
            _apply_banner_style(p, shd_banner_xml)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.page_break_before = break_before_part2

        elif "PART 3" in text_upper or "CHECK YOUR UNDERSTANDING" in text_upper:
            current_section = "PART3"
            _apply_banner_style(p, shd_banner_xml)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.page_break_before = (not break_before_part2)

        elif p == doc.paragraphs[0]:
            # Title
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)

        elif re.match(r"^Tool\s*\d+", text, re.IGNORECASE):
            # Tool subheadings (close to its table)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)

        elif "identify the mistake" in text.lower():
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)

        elif current_section == "PART3" and text:
            # CYU Questions
            p.paragraph_format.space_before = cyu_space_before
            p.paragraph_format.space_after = cyu_space_after

    # -------------------------------------------------------------
    # 3. FORMAT TABLES (WIDTHS, BORDERS, SHADING, PADDING)
    # -------------------------------------------------------------
    borders_xml = (
        f'<w:tblBorders {w_ns}>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:left w:val="single" w:sz="8" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:right w:val="single" w:sz="8" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:insideH w:val="single" w:sz="6" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'<w:insideV w:val="single" w:sz="6" w:space="0" w:color="{BORDER_COLOR}"/>'
        f'</w:tblBorders>'
    )

    cell_mar_xml = (
        f'<w:tblCellMar {w_ns}>'
        f'<w:top w:w="140" w:type="dxa"/>'
        f'<w:bottom w:w="140" w:type="dxa"/>'
        f'<w:left w:w="180" w:type="dxa"/>'
        f'<w:right w:w="180" w:type="dxa"/>'
        f'</w:tblCellMar>'
    )

    cant_split_xml = f'<w:cantSplit {w_ns}/>'
    tbl_header_xml = f'<w:tblHeader {w_ns}/>'
    valign_center_xml = f'<w:vAlign {w_ns} w:val="center"/>'
    shd_header_xml = f'<w:shd {w_ns} w:fill="{HEADER_BG}"/>'
    jc_left_xml = f'<w:jc {w_ns} w:val="left"/>'

    for table in doc.tables:
        table.autofit = False
        tblPr = table._tbl.tblPr

        # Apply borders
        existing_borders = tblPr.find(qn("w:tblBorders"))
        if existing_borders is not None:
            tblPr.remove(existing_borders)
        tblPr.append(parse_xml(borders_xml))

        # Apply cell padding
        existing_mar = tblPr.find(qn("w:tblCellMar"))
        if existing_mar is not None:
            tblPr.remove(existing_mar)
        tblPr.append(parse_xml(cell_mar_xml))

        # Left alignment
        existing_jc = tblPr.find(qn("w:jc"))
        if existing_jc is not None:
            tblPr.remove(existing_jc)
        tblPr.append(parse_xml(jc_left_xml))

        # Determine table column widths
        first_row_text = " ".join([c.text.lower() for c in table.rows[0].cells])
        is_fix_mistake = ("correct" in first_row_text or "mistake" in first_row_text)

        if len(table.columns) == 2:
            if is_fix_mistake:
                col_widths = [4.0, 2.5]       # 6.5" full width
                min_body_height = Inches(0.85)
            else:
                col_widths = [3.2, 1.8]       # 5.0" compact Tool width (~2/3 and ~1/3)
                min_body_height = Inches(0.95)
        else:
            col_widths = [6.5 / len(table.columns)] * len(table.columns)
            min_body_height = Inches(0.7)

        total_width_dxa = int(sum(col_widths) * 1440)
        tblW_xml = f'<w:tblW {w_ns} w:w="{total_width_dxa}" w:type="dxa"/>'
        existing_tblW = tblPr.find(qn("w:tblW"))
        if existing_tblW is not None:
            tblPr.remove(existing_tblW)
        tblPr.append(parse_xml(tblW_xml))

        for r_idx, row in enumerate(table.rows):
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(cant_split_xml))

            is_header = (r_idx == 0)
            if is_header:
                trPr.append(parse_xml(tbl_header_xml))
            else:
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                row.height = min_body_height

            for c_idx, width_in in enumerate(col_widths):
                if c_idx < len(row.cells):
                    cell = row.cells[c_idx]
                    cell.width = Inches(width_in)
                    tcPr = cell._tc.get_or_add_tcPr()

                    cell_width_dxa = int(width_in * 1440)
                    tcW_xml = f'<w:tcW {w_ns} w:w="{cell_width_dxa}" w:type="dxa"/>'
                    existing_tcW = tcPr.find(qn("w:tcW"))
                    if existing_tcW is not None:
                        tcPr.remove(existing_tcW)
                    tcPr.append(parse_xml(tcW_xml))

                    tcPr.append(parse_xml(valign_center_xml))

                    if is_header:
                        existing_shd = tcPr.find(qn("w:shd"))
                        if existing_shd is not None:
                            tcPr.remove(existing_shd)
                        tcPr.append(parse_xml(shd_header_xml))

                        for p in cell.paragraphs:
                            p.paragraph_format.space_before = Pt(4)
                            p.paragraph_format.space_after = Pt(4)
                            for run in p.runs:
                                run.bold = True
                    else:
                        for p in cell.paragraphs:
                            if p.paragraph_format.space_before is None or p.paragraph_format.space_before < Pt(2):
                                p.paragraph_format.space_before = Pt(3)
                            if p.paragraph_format.space_after is None or p.paragraph_format.space_after < Pt(2):
                                p.paragraph_format.space_after = Pt(3)

    try:
        doc.save(docx_path)
        print(f"Document formatted successfully with {tool_count} tools detected.")
    except Exception as e:
        print(f"Warning: Could not save formatted document: {e}")


def _apply_banner_style(p, shd_xml):
    """Applies a dark charcoal background banner with bold white text."""
    pPr = p._element.get_or_add_pPr()
    existing_shd = pPr.find(qn("w:shd"))
    if existing_shd is not None:
        pPr.remove(existing_shd)
    pPr.append(parse_xml(shd_xml))

    for run in p.runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.bold = True
        run.font.size = Pt(11)
