import io
import os
import base64
from typing import Tuple, Optional

# Supported extensions
TEXT_PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}


def extract_docx_text(file_stream) -> str:
    """Extract paragraphs and tables from a Word (.docx) document."""
    try:
        import docx
        doc = docx.Document(file_stream)
        text_lines = []

        # Extract paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                text_lines.append(text)

        # Extract tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(row_data):
                    table_rows.append(" | ".join(row_data))
            if table_rows:
                text_lines.append("\n[Table Content]\n" + "\n".join(table_rows))

        return "\n\n".join(text_lines)
    except Exception as e:
        return f"[Error parsing DOCX: {str(e)}]"


def extract_pdf_text(file_stream) -> Tuple[str, bool]:
    """
    Extract text from a PDF stream using pypdf or pdfplumber.
    Returns (extracted_text, is_scanned_or_empty).
    """
    text = ""
    # Try pypdf first
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_stream)
        extracted = []
        for page in reader.pages:
            p_text = page.extract_text() or ""
            if p_text.strip():
                extracted.append(p_text.strip())
        text = "\n\n".join(extracted)
    except ImportError:
        try:
            import pdfplumber
            with pdfplumber.open(file_stream) as pdf:
                extracted = []
                for page in pdf.pages:
                    p_text = page.extract_text() or ""
                    if p_text.strip():
                        extracted.append(p_text.strip())
                text = "\n\n".join(extracted)
        except Exception as e:
            return f"[Error reading PDF: {str(e)}]", True
    except Exception as e:
        return f"[Error reading PDF: {str(e)}]", True

    # If minimal or no text extracted, likely a scanned PDF
    is_scanned = len(text.strip()) < 50
    return text, is_scanned


def encode_image_base64(file_stream) -> str:
    """Encodes file stream bytes to base64 string."""
    file_bytes = file_stream.read()
    if hasattr(file_stream, "seek"):
        file_stream.seek(0)
    return base64.b64encode(file_bytes).decode("utf-8")
